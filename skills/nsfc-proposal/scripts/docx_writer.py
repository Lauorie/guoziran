"""用 python-docx 把块写成 Word（供 export_docx.py 调用）。版式：A4，正文宋体/Times New Roman 小四，1.25 倍行距，
首行缩进两字；标题黑体；图居中，图注、表注五号；引用编号上标；【待补充：…】占位黄色高亮，方便申请人逐个补。"""

from __future__ import annotations

import os
import re
from pathlib import Path
from typing import Any

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_BREAK, WD_COLOR_INDEX, WD_PARAGRAPH_ALIGNMENT
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from docx_bib import cite_label
from docx_latex import Block, Run, _text_run

BODY_FONT, HEAD_FONT, LATIN_FONT = "宋体", "黑体", "Times New Roman"
HEADING_SIZES = {1: 14, 2: 12, 3: 12, 4: 12}
MAX_IMAGE_CM = 14.0
MAX_IMAGE_PX = 1800      # 嵌入前把过宽的图缩到这个宽度，免得 2K 图让文档有十几 MB
LINE_SPACING = 1.25


def _fonts(run, east: str = BODY_FONT, size: float | None = None, bold: bool | None = None) -> None:
    run.font.name = LATIN_FONT
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = rpr.makeelement(qn("w:rFonts"), {})
        rpr.append(rfonts)
    rfonts.set(qn("w:eastAsia"), east)
    if size:
        run.font.size = Pt(size)
    if bold is not None:
        run.font.bold = bold


def _setup(doc) -> None:
    section = doc.sections[0]
    section.page_width, section.page_height = Cm(21.0), Cm(29.7)
    section.top_margin = section.bottom_margin = Cm(2.5)
    section.left_margin = section.right_margin = Cm(3.0)
    normal = doc.styles["Normal"]
    normal.font.name, normal.font.size = LATIN_FONT, Pt(12)
    normal.element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), BODY_FONT)
    normal.paragraph_format.line_spacing = LINE_SPACING
    normal.paragraph_format.space_after = Pt(0)
    for level, size in HEADING_SIZES.items():
        style = doc.styles[f"Heading {level}"]
        style.font.name, style.font.size, style.font.bold = LATIN_FONT, Pt(size), True
        style.font.italic = False
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), HEAD_FONT)
        style.paragraph_format.space_before, style.paragraph_format.space_after = Pt(6), Pt(6)


def _add_runs(paragraph, runs: list[Run], ctx: dict[str, Any], size: float | None = None) -> int:
    placeholders = 0
    for run in runs:
        if run.cite:
            piece = paragraph.add_run(cite_label(run.cite, ctx["cites"]))
            piece.font.superscript = ctx["superscript"]
        elif run.ref is not None:
            label = ctx["labels"].get(run.ref)
            if label is None:
                ctx["warnings"].append(f"找不到图表标签 {run.ref}")
            piece = paragraph.add_run(label or "?")
        else:
            piece = paragraph.add_run(run.text)
            if run.placeholder:
                piece.font.highlight_color = WD_COLOR_INDEX.YELLOW
                placeholders += 1
        _fonts(piece, size=size)
        piece.font.bold = run.bold or None
        piece.font.italic = run.italic or None
        piece.font.underline = run.underline or None
    return placeholders


def _body_paragraph(doc, runs: list[Run], ctx: dict[str, Any], indent: bool = True):
    paragraph = doc.add_paragraph()
    if indent:
        paragraph.paragraph_format.first_line_indent = Pt(24)
    ctx["placeholders"] += _add_runs(paragraph, runs, ctx)
    return paragraph


def _caption(doc, prefix: str, runs: list[Run], ctx: dict[str, Any]) -> None:
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
    head = paragraph.add_run(prefix + "　")
    _fonts(head, size=10.5, bold=True)
    ctx["placeholders"] += _add_runs(paragraph, runs, ctx, size=10.5)


def _figure(doc, block: Block, ctx: dict[str, Any]) -> None:
    path = (ctx["root"] / (block.path or "")).resolve()
    if path.suffix.lower() not in {".png", ".jpg", ".jpeg", ".gif", ".bmp"} and path.with_suffix(".png").is_file():
        path = path.with_suffix(".png")
    if ctx["root"].resolve() in path.parents and path.is_file() and path.suffix.lower() in {".png", ".jpg", ".jpeg", ".gif", ".bmp"}:
        paragraph = doc.add_paragraph()
        paragraph.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
        paragraph.add_run().add_picture(_shrunk(path, ctx), width=Cm(MAX_IMAGE_CM))
        ctx["figures"] += 1
    else:
        ctx["warnings"].append(f"图片不存在或不是位图，未插入：{block.path}")
        _body_paragraph(doc, [Run(text=f"【待补充：图片 {block.path}】", placeholder=True)], ctx, indent=False)
    if block.text:
        _caption(doc, f"图 {block.text}", block.runs, ctx)


def _shrunk(path: Path, ctx: dict[str, Any]):
    """宽于 MAX_IMAGE_PX 的图缩小后以内存流嵌入；没有 Pillow 时用原图。"""
    try:
        from io import BytesIO

        from PIL import Image
    except ImportError:
        return str(path)
    try:
        with Image.open(path) as image:
            if image.width <= MAX_IMAGE_PX:
                return str(path)
            height = round(image.height * MAX_IMAGE_PX / image.width)
            small = image.convert("RGB" if image.mode not in ("RGB", "L") else image.mode).resize((MAX_IMAGE_PX, height))
            stream = BytesIO()
            small.save(stream, format="PNG", optimize=True)
            stream.seek(0)
            return stream
    except OSError:
        ctx["warnings"].append(f"图片读不了，按原文件嵌入：{path.name}")
        return str(path)


def _table(doc, block: Block, ctx: dict[str, Any]) -> None:
    if block.text:
        _caption(doc, f"表 {block.text}", block.runs, ctx)
    width = max((len(row) for row in block.rows), default=0)
    if not width:
        return
    table = doc.add_table(rows=len(block.rows), cols=width)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for r, row in enumerate(block.rows):
        for c in range(width):
            cell = table.cell(r, c)
            cell.paragraphs[0].text = ""
            ctx["placeholders"] += _add_runs(cell.paragraphs[0], row[c] if c < len(row) else [], ctx, size=10.5)
    ctx["tables"] += 1


def _references(doc, references: list[tuple[int, str]], ctx: dict[str, Any]) -> None:
    title = doc.add_paragraph()
    _fonts(title.add_run("参考文献"), east=HEAD_FONT, size=12, bold=True)
    for number, text in references:
        paragraph = doc.add_paragraph()
        paragraph.paragraph_format.left_indent = Pt(24)
        paragraph.paragraph_format.first_line_indent = Pt(-24)
        paragraph.paragraph_format.line_spacing = 1.0
        piece = paragraph.add_run(f"[{number}] {text}")
        _fonts(piece, size=10.5)
        if text.startswith("【待补充"):
            piece.font.highlight_color = WD_COLOR_INDEX.YELLOW
    ctx["references_written"] = True


def write(target: Path, blocks: list[Block], *, cites: dict[str, int], labels: dict[str, str],
          references: list[tuple[int, str]], appendix: list, root: Path, superscript: bool,
          warnings: list[str]) -> dict[str, int]:
    doc = Document()
    _setup(doc)
    ctx: dict[str, Any] = {"cites": cites, "labels": labels, "superscript": superscript, "warnings": warnings,
                           "root": root, "placeholders": 0, "figures": 0, "tables": 0, "references_written": False}
    title = doc.add_paragraph()
    title.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
    _fonts(title.add_run("报告正文"), east=HEAD_FONT, size=16, bold=True)
    paragraphs = 0
    for block in blocks:
        if block.kind == "heading":
            paragraph = doc.add_paragraph(style=f"Heading {min(block.level, 4)}")
            _add_runs(paragraph, block.runs, ctx)
        elif block.kind == "para":
            _body_paragraph(doc, block.runs, ctx)
            paragraphs += 1
        elif block.kind == "list":
            for n, item in enumerate(block.items, 1):
                prefix = Run(text=f"（{n}）" if block.ordered else "• ")
                _body_paragraph(doc, [prefix] + item, ctx)
        elif block.kind == "figure":
            _figure(doc, block, ctx)
        elif block.kind == "table":
            _table(doc, block, ctx)
        elif block.kind == "math":
            paragraph = doc.add_paragraph()
            paragraph.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
            piece = paragraph.add_run(block.text)
            _fonts(piece)
            piece.font.italic = True
        elif block.kind == "bibliography" and references and not ctx["references_written"]:
            _references(doc, references, ctx)
    if references and not ctx["references_written"]:
        _references(doc, references, ctx)
    for heading, sections in appendix:
        doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
        doc.add_paragraph(style="Heading 1").add_run(heading)
        for sub, lines in sections:
            if sub:
                doc.add_paragraph(style="Heading 2").add_run(sub)
            for line in lines:
                runs = [r for part in _split_bold(line) for r in _text_run(part.strip("*"), {"bold": part.startswith("**")})]
                _body_paragraph(doc, runs, ctx)
    target.parent.mkdir(parents=True, exist_ok=True)
    tmp = target.with_name(target.name + ".tmp")
    doc.save(str(tmp))
    os.replace(tmp, target)
    return {"paragraphs": paragraphs, "figures": ctx["figures"], "tables": ctx["tables"],
            "placeholders": ctx["placeholders"]}


def _split_bold(line: str) -> list[str]:
    """Markdown 行里的 **加粗** 拆成单独的片段。"""
    return [p for p in re.split(r"(\*\*[^*]+\*\*)", line) if p]
