# tests/nsfc-proposal/test_np_export_docx.py
"""export_docx.py：NSFC 正文项目 → 一个 Word 文档（正文 + 参考文献 + 附题目与摘要）。不联网，不编译 LaTeX。"""
from __future__ import annotations

import base64
import json
import os
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path

ROOT = Path(__file__).resolve().parents[2]
SCRIPTS = ROOT / "skills/nsfc-proposal/scripts"
sys.path.insert(0, str(SCRIPTS))

try:
    import docx  # noqa: F401
    HAVE_DOCX = True
except ImportError:  # pragma: no cover
    HAVE_DOCX = False

from docx_bib import cite_label, format_authors, parse_bib, reference_text  # noqa: E402
from docx_blocks import parse_chapter  # noqa: E402
from docx_latex import inline, math_text  # noqa: E402

PNG_1X1 = base64.b64decode(
    "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJAAAADUlEQVR42mNkYPhfDwAChwGA60e6kgAAAABJRU5ErkJggg==")
BIB = """@article{zhu2026,
  title = {{垃圾分类政策前后居民参与意愿对比研究}},
  author = {朱姝娲 and 张俊杰 and 高莉洁 and 肖黎姗},
  journal = {环境卫生工程}, year = {2026}, volume = {34}, number = {3}, pages = {112--119},
}
@inproceedings{kojima2022,
  title = {Large Language Models Are Zero-Shot Reasoners},
  author = {Kojima, Takeshi and Gu, Shixiang (Shane)},
  booktitle = {Advances in Neural Information Processing Systems 35}, year = {2022}, pages = {22199--22213},
  doi = {10.52202/068431-1613},
}
"""
JUSTIFICATION = r"""% 1.1.立项依据
\subsubsection{研究背景}
居民分类正确率在宣传期后回落\cite{zhu2026}，零样本推理已有报道\cite{kojima2022,zhu2026}。
本项目的框架见图~\ref{fig:framework}，参数见表~\ref{tab:params}。

\paragraph{研究缺口} 现有研究多测量意愿。【待补充：前期调查的样本量】
\begin{figure}[htbp]
  \centering
  \includegraphics[width=0.9\textwidth]{figures/nsfc-proposal/fig1.png}
  \caption{研究框架}\label{fig:framework}
\end{figure}
\begin{table}[htbp]
  \centering
  \caption{主要参数}\label{tab:params}
  \begin{tabular}{p{3cm}l}
    \toprule 指标 & 数值 \\ \midrule \textbf{样本量} & $n = 45$ \\ \bottomrule
  \end{tabular}
\end{table}
\begin{enumerate}
  \item 第一项任务；
  \item 第二项任务，引用未入库的文献\cite{missing2020}。
\end{enumerate}
"""
MAIN_TEX = r"""\documentclass{ctexart}
\begin{document}
\section{（一）立项依据}
\subsection{（为什么要开展此项研究，研究的科学技术价值如何）}
\input{extraTex/1.1.立项依据.tex}
\input{references/reference.tex}
\section{（二）研究内容}
\input{extraTex/2.1.研究内容.tex}
\end{document}
"""
ABSTRACTS = """# 标题建议

1. 城市居民垃圾分类习惯的形成机制研究

# 中文摘要

本项目研究……【待补充：前期研究的主要发现】

# English Abstract

This project studies **habit formation**.
"""


class InlineTest(unittest.TestCase):
    def test_inline_commands(self) -> None:
        runs = inline(r"\textbf{加粗}与\emph{斜体}，转义 50\%、A\&B，引用\cite{a, b}，图~\ref{fig:x}--完")
        self.assertEqual([r.text for r in runs if r.bold], ["加粗"])
        self.assertEqual([r.text for r in runs if r.italic], ["斜体"])
        self.assertEqual([r.cite for r in runs if r.cite], [("a", "b")])
        self.assertEqual([r.ref for r in runs if r.ref is not None], ["fig:x"])
        text = "".join(r.text for r in runs)
        self.assertIn("50%", text)
        self.assertIn("A&B", text)
        self.assertIn("–完", text)

    def test_math_and_unknown_commands(self) -> None:
        self.assertEqual(math_text(r"\alpha = 0.05, R^2 \geq 0.8, \frac{a}{b}"), "α = 0.05, R² ≥ 0.8, a/b")
        runs = inline(r"\NSFCCustom{保留文字}\noindent 后文")
        self.assertEqual("".join(r.text for r in runs), "保留文字后文")

    def test_template_heading_macros(self) -> None:
        text = r"{\punctstyle{banjiao}1. \textbf{研究基础与可行性分析}（与本项目相关的研究工作积累和已\linebreak{}取得的研究工作成绩）；}"
        self.assertEqual("".join(r.text for r in inline(text)), "1. 研究基础与可行性分析（与本项目相关的研究工作积累和已取得的研究工作成绩）；")
        text = r"\texorpdfstring{\hspace{0.1em}}{}（应提炼出特\linebreak{}色与创新点）"
        self.assertEqual("".join(r.text for r in inline(text)), "（应提炼出特色与创新点）")

    def test_cjk_line_breaks_are_joined(self) -> None:
        self.assertEqual("".join(r.text for r in inline("第一行\n第二行 English\nword")), "第一行第二行 English word")

    def test_blocks(self) -> None:
        blocks = parse_chapter(JUSTIFICATION)
        kinds = [b.kind for b in blocks]
        self.assertEqual(kinds[0], "heading")
        self.assertIn("figure", kinds)
        self.assertIn("table", kinds)
        self.assertIn("list", kinds)
        para = next(b for b in blocks if b.kind == "para" and b.runs and b.runs[0].bold)
        self.assertEqual(para.runs[0].text, "研究缺口")  # \paragraph 是段首加粗
        self.assertTrue(any(r.placeholder for r in para.runs))
        table = next(b for b in blocks if b.kind == "table")
        self.assertEqual([[''.join(r.text for r in c) for c in row] for row in table.rows], [["指标", "数值"], ["样本量", "n = 45"]])


class BibTest(unittest.TestCase):
    def test_gbt7714(self) -> None:
        entries = parse_bib(BIB + "\n@comment{ignored}\n")
        self.assertEqual(reference_text(entries["zhu2026"]),
                         "朱姝娲, 张俊杰, 高莉洁, 等. 垃圾分类政策前后居民参与意愿对比研究[J]. 环境卫生工程, 2026, 34(3): 112-119.")
        self.assertEqual(reference_text(entries["kojima2022"]),
                         "Kojima T, Gu S. Large Language Models Are Zero-Shot Reasoners[C]//Advances in Neural Information "
                         "Processing Systems 35. 2022: 22199-22213. DOI: 10.52202/068431-1613.")
        self.assertEqual(format_authors("Smith, John and Doe, Jane Q. and A, B and C, D"), "Smith J, Doe J Q, A B, et al")

    def test_preprint_is_online_resource(self) -> None:
        entries = parse_bib("""@misc{cho2025breaking,
  title = {{Breaking the Safety-Capability Tradeoff}},
  author = {Cho, Dongkyu Derek and Song, Huan},
  howpublished = {arXiv preprint arXiv:2511.21050},
  year = {2025},
  doi = {10.48550/arxiv.2511.21050},
  url = {https://arxiv.org/abs/2511.21050},
}
@misc{lee2024open,
  title = {Open Test Notes},
  author = {Lee, Ann},
  year = {2024},
  url = {https://example.org/notes},
}
""")
        self.assertEqual(reference_text(entries["cho2025breaking"]),
                         "Cho D D, Song H. Breaking the Safety-Capability Tradeoff[EB/OL]. arXiv preprint arXiv:2511.21050, 2025. "
                         "DOI: 10.48550/arxiv.2511.21050.")
        self.assertEqual(reference_text(entries["lee2024open"]), "Lee A. Open Test Notes[EB/OL]. 2024. https://example.org/notes.")

    def test_cite_label_compresses_ranges(self) -> None:
        cites = {"a": 1, "b": 2, "c": 3, "d": 5}
        self.assertEqual(cite_label(["c", "a", "b"], cites), "[1-3]")
        self.assertEqual(cite_label(["a", "b"], cites), "[1,2]")
        self.assertEqual(cite_label(["a", "d"], cites), "[1,5]")
        self.assertEqual(cite_label(["zzz"], cites), "[?]")


@unittest.skipUnless(HAVE_DOCX, "本机没有 python-docx")
class ExportTest(unittest.TestCase):
    def setUp(self) -> None:
        tmp = tempfile.TemporaryDirectory()
        self.addCleanup(tmp.cleanup)
        self.root = Path(tmp.name) / "nsfc-project"
        (self.root / "extraTex").mkdir(parents=True)
        (self.root / "references").mkdir()
        (self.root / "main.tex").write_text(MAIN_TEX, encoding="utf-8")
        (self.root / "references" / "reference.tex").write_text(
            "\\bibliographystyle{gbt7714-numerical}\n\\bibliography{references/myexample}\n", encoding="utf-8")
        (self.root / "extraTex" / "2.1.研究内容.tex").write_text("研究内容正文。\n", encoding="utf-8")
        (self.root / "extraTex" / "1.1.立项依据.tex").write_text(JUSTIFICATION, encoding="utf-8")
        (self.root / "references" / "myexample.bib").write_text(BIB, encoding="utf-8")
        (self.root / "figures" / "nsfc-proposal").mkdir(parents=True)
        (self.root / "figures" / "nsfc-proposal" / "fig1.png").write_bytes(PNG_1X1)
        (self.root / "NSFC-ABSTRACTS.md").write_text(ABSTRACTS, encoding="utf-8")

    def export(self) -> dict:
        import export_docx
        return export_docx.export(self.root, config={"export": {}})

    def test_document_structure(self) -> None:
        from docx import Document
        summary = self.export()
        self.assertEqual((summary["citations"], summary["figures"], summary["tables"]), (3, 1, 1))
        self.assertTrue(any("missing2020" in w for w in summary["warnings"]))
        doc = Document(summary["docx"])
        texts = [p.text for p in doc.paragraphs]
        self.assertEqual(texts[0], "报告正文")
        self.assertIn("（一）立项依据", texts)
        self.assertIn("研究背景", texts)
        body = next(t for t in texts if t.startswith("居民分类正确率"))
        self.assertIn("回落[1]", body)
        self.assertIn("[1,2]", body)  # 首次出现的顺序：zhu2026=1，kojima2022=2
        self.assertIn("见图 1", body.replace(" ", " "))
        self.assertIn("见表 1", body.replace(" ", " "))
        refs_at = texts.index("参考文献")
        self.assertTrue(texts[refs_at + 1].startswith("[1] 朱姝娲"))
        self.assertTrue(texts[refs_at + 2].startswith("[2] Kojima T"))
        self.assertIn("【待补充：参考文献 missing2020 不在 .bib 中】", texts[refs_at + 3])
        self.assertLess(refs_at, texts.index("（二）研究内容"))  # 参考文献在立项依据之后
        self.assertIn("图 1　研究框架", texts)
        self.assertIn("表 1　主要参数", texts)
        self.assertIn("附：题目与摘要", texts)
        self.assertIn("This project studies habit formation.", texts)
        self.assertEqual(len(doc.inline_shapes), 1)
        self.assertEqual([[c.text for c in row.cells] for row in doc.tables[0].rows], [["指标", "数值"], ["样本量", "n = 45"]])
        superscripts = [r.text for p in doc.paragraphs for r in p.runs if r.font.superscript]
        self.assertIn("[1]", superscripts)
        highlighted = [r.text for p in doc.paragraphs for r in p.runs if r.font.highlight_color is not None]
        self.assertIn("【待补充：前期调查的样本量】", highlighted)
        self.assertIn("【待补充：前期研究的主要发现】", highlighted)

    def test_missing_image_becomes_placeholder(self) -> None:
        (self.root / "figures" / "nsfc-proposal" / "fig1.png").unlink()
        summary = self.export()
        self.assertEqual(summary["figures"], 0)
        self.assertTrue(any("图片不存在" in w for w in summary["warnings"]))

    def test_cli(self) -> None:
        env = dict(os.environ, PYTHONPYCACHEPREFIX=str(Path(tempfile.gettempdir()) / "nsfc-proposal-pycache"))
        proc = subprocess.run([sys.executable, str(SCRIPTS / "export_docx.py"), "--project-root", str(self.root)],
                              env=env, capture_output=True, text=True, timeout=120)
        self.assertEqual(proc.returncode, 0, proc.stderr)
        summary = json.loads(proc.stdout)
        self.assertTrue(Path(summary["docx"]).is_file())
        self.assertEqual(Path(summary["docx"]).name, "申请书正文.docx")
        proc = subprocess.run([sys.executable, str(SCRIPTS / "export_docx.py"), "--project-root", str(self.root / "nope")],
                              env=env, capture_output=True, text=True, timeout=60)
        self.assertEqual(proc.returncode, 2)


if __name__ == "__main__":
    unittest.main()
